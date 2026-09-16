"""Full API workflow test: upload -> translate -> review -> edit -> commit -> download.

Uses TestClient plus a FakeChain (no real LLM, no browser). The goal is the
wiring between layers, not UI behavior.
"""

from __future__ import annotations

import io
import time

import pytest
from fastapi.testclient import TestClient
from pptx import Presentation
from pptx.util import Inches

from api import app
from src.chains.translation_chain import TranslationOutput
from src.services.job_manager import get_job_manager
from src.utils.config import get_settings


class FakeChain:
    """Deterministic stand-in for the LangChain runnable (never touches network)."""

    @staticmethod
    def _translate(batch) -> TranslationOutput:
        return TranslationOutput(
            translations=[
                f"가짜번역:{paragraph.original_text}"
                for paragraph in batch.get("paragraphs", [])
            ]
        )

    def batch_as_completed(self, submitted, config=None):
        for local_idx, batch in enumerate(submitted):
            yield local_idx, self._translate(batch)

    def invoke(self, batch, config=None):
        return self._translate(batch)


def _make_deck() -> bytes:
    """Two plain single-color text boxes (no color pass, no LLM outside FakeChain)."""
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1))
    box.text_frame.text = "Hello world"
    box2 = slide.shapes.add_textbox(Inches(1), Inches(3), Inches(5), Inches(1))
    box2.text_frame.text = "Second line"
    buffer = io.BytesIO()
    presentation.save(buffer)
    return buffer.getvalue()


@pytest.fixture
def workflow_env(monkeypatch, tmp_path):
    """Dummy API key + isolated quality ledger (get_settings is lru-cached)."""
    monkeypatch.setenv("ANTHROPIC_API_KEY", "test")
    monkeypatch.setenv("TRANSLATION_QUALITY_DIR", str(tmp_path))
    get_settings.cache_clear()
    yield
    get_settings.cache_clear()


def _wait_for_completion(client: TestClient, job_id: str, timeout: float = 30.0) -> None:
    deadline = time.monotonic() + timeout
    while time.monotonic() < deadline:
        status = client.get(f"/api/v1/jobs/{job_id}")
        assert status.status_code == 200, status.text
        state = status.json()["state"]
        if state == "completed":
            return
        assert state not in ("failed", "cancelled"), status.json()
        time.sleep(0.1)
    raise AssertionError(f"job {job_id} did not complete within {timeout}s")


def test_upload_translate_review_edit_commit_download(monkeypatch, workflow_env):
    """One pass through the whole translation pipeline via HTTP."""
    monkeypatch.setattr(
        "src.services.translation_service.create_translation_chain",
        lambda **kwargs: FakeChain(),
    )
    manager = get_job_manager()

    with TestClient(app) as client:
        response = client.post(
            "/api/v1/jobs",
            files={
                "ppt_file": (
                    "workflow.pptx",
                    _make_deck(),
                    "application/octet-stream",
                )
            },
            data={
                "provider": "anthropic",
                "model": "claude-sonnet-5",
                "source_lang": "영어",
                "target_lang": "한국어",
            },
        )
        assert response.status_code == 200, response.text
        job_id = response.json()["job_id"]
        try:
            _wait_for_completion(client, job_id)

            fragments = client.get(f"/api/v1/jobs/{job_id}/fragments")
            assert fragments.status_code == 200, fragments.text
            body = fragments.json()
            assert body["total"] == 2
            assert [f["target"] for f in body["fragments"]] == [
                "가짜번역:Hello world",
                "가짜번역:Second line",
            ]

            edit = client.post(
                f"/api/v1/jobs/{job_id}/review/block",
                json={
                    "edits": {"0": "첫 번째 수정", "1": "두 번째 수정"},
                    "expected_revision": body["revision"],
                },
            )
            assert edit.status_code == 200, edit.text
            assert edit.json()["revision"] == body["revision"] + 1

            commit = client.post(
                f"/api/v1/jobs/{job_id}/review/commit",
                json={"expected_revision": body["revision"] + 1},
            )
            assert commit.status_code == 200, commit.text
            assert commit.json()["dirty"] is False

            download = client.get(f"/api/v1/jobs/{job_id}/result")
            assert download.status_code == 200
            rendered = Presentation(io.BytesIO(download.content))
            texts = [
                shape.text
                for shape in rendered.slides[0].shapes
                if shape.has_text_frame
            ]
            assert "첫 번째 수정" in texts
            assert "두 번째 수정" in texts
        finally:
            manager._jobs.pop(job_id, None)
