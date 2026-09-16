"""Tests for ADR-0006 scope accounting in src/core/ppt_parser.py.

Chart/SmartArt/OLE shapes are excluded from translation. The parser must
count them so the warning is accurate — including SmartArt, whose
shape_type is None (detected by diagram graphicData uri instead).
"""

from __future__ import annotations

import io
import types

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches

from src.core.ppt_parser import PPTParser

_DIAGRAM_URI = "http://schemas.openxmlformats.org/drawingml/2006/diagram"


def _chart_deck() -> io.BytesIO:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(4), Inches(1)).text_frame.text = (
        "Hello"
    )
    cd = CategoryChartData()
    cd.categories = ["Q1", "Q2"]
    cd.add_series("Sales", (1, 2))
    slide.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(0.5), Inches(2), Inches(5), Inches(3), cd
    )
    buf = io.BytesIO()
    prs.save(buf)
    return buf


def test_chart_excluded_from_paragraphs_but_counted():
    paragraphs, presentation = PPTParser().extract_paragraphs(_chart_deck())
    assert [p.original_text for p in paragraphs] == ["Hello"]
    assert PPTParser.count_skipped_shapes(presentation) == 1


def test_smartart_with_none_shape_type_is_counted():
    fake = types.SimpleNamespace(
        shape_type=None,
        _element=types.SimpleNamespace(graphicData_uri=_DIAGRAM_URI),
    )
    presentation = types.SimpleNamespace(
        slides=[types.SimpleNamespace(shapes=[fake])]
    )
    assert PPTParser.count_skipped_shapes(presentation) == 1


def test_plain_text_deck_counts_zero():
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(4), Inches(1)).text_frame.text = (
        "Hello"
    )
    buf = io.BytesIO()
    prs.save(buf)
    _, presentation = PPTParser().extract_paragraphs(buf)
    assert PPTParser.count_skipped_shapes(presentation) == 0
