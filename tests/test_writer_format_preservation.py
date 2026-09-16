"""Regression tests for hyperlink and manual line-break preservation.

Covers two writer bugs where ``run.text = ...`` was used directly:

- A hyperlinked run merged with adjacent plain-text runs (``_rpr_key``
  ignored ``<a:hlinkClick>``), turning the whole paragraph into one link.
- Manual line breaks (``<a:br/>``) mishandled: the setter only replaces
  the first ``<a:t>``, so stale siblings leaked source-language text and
  raw ``\\n`` never rendered as a break. Extraction dropped post-break
  text as well (``run.text`` reads the first ``<a:t>`` only).
"""

from __future__ import annotations

import io
import types
import unittest

from lxml import etree
from pptx import Presentation
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml.ns import qn
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches

from src.core.ppt_parser import PPTParser
from src.core.ppt_writer import PPTWriter, _group_runs_by_format
from src.utils.helpers import run_text_with_breaks

_A = "http://schemas.openxmlformats.org/drawingml/2006/main"


def _add_paragraph(prs: Presentation):
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank layout
    shape = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
    paragraph = shape.text_frame.paragraphs[0]
    return slide, paragraph


def _append_break(run, trailing_text: str = "") -> None:
    """Append a manual line break (Shift+Enter) plus optional text."""
    run._r.append(OxmlElement("a:br"))
    if trailing_text:
        t = OxmlElement("a:t")
        t.text = trailing_text
        run._r.append(t)


def _append_hyperlink(run, slide, url: str = "https://example.com") -> None:
    """Attach an external hyperlink to a run."""
    r_id = slide.part.relate_to(url, RT.HYPERLINK, is_external=True)
    hlink = OxmlElement("a:hlinkClick")
    hlink.set(qn("r:id"), r_id)
    run._r.append(hlink)


def _run_xml(run) -> str:
    return etree.tostring(run._r, encoding="unicode")


class RunTextWithBreaksTestCase(unittest.TestCase):
    def test_full_text_includes_post_break_content(self):
        prs = Presentation()
        _, paragraph = _add_paragraph(prs)
        run = paragraph.add_run()
        run.text = "Hello"
        _append_break(run, "World")

        self.assertEqual(run_text_with_breaks(run), "Hello\nWorld")
        # python-pptx's own getter only sees the first <a:t>
        self.assertEqual(run.text, "Hello")


class ExtractManualBreakTestCase(unittest.TestCase):
    def test_parser_keeps_text_after_manual_break(self):
        prs = Presentation()
        _, paragraph = _add_paragraph(prs)
        run = paragraph.add_run()
        run.text = "Hello"
        _append_break(run, "World")

        buf = io.BytesIO()
        prs.save(buf)
        paragraphs, _ = PPTParser().extract_paragraphs(buf)

        self.assertEqual(len(paragraphs), 1)
        self.assertEqual(paragraphs[0].original_text, "Hello\nWorld")


class WriterManualBreakTestCase(unittest.TestCase):
    def _apply(self, paragraph, original_text: str, translation: str):
        info = types.SimpleNamespace(
            paragraph=paragraph,
            original_text=original_text,
            is_note=False,
            slide_index=0,
            shape_index=0,
            paragraph_index=0,
        )
        PPTWriter().apply_translations(
            [info], [translation], Presentation(), text_fit_mode="none"
        )

    def test_newlines_become_br_and_stale_siblings_removed(self):
        prs = Presentation()
        _, paragraph = _add_paragraph(prs)
        run = paragraph.add_run()
        run.text = "Hello"
        _append_break(run, "World")

        self._apply(paragraph, "Hello\nWorld", "안녕\n세계")

        xml = _run_xml(paragraph.runs[0])
        self.assertNotIn("Hello", xml)
        self.assertNotIn("World", xml)
        texts = [
            el.text or ""
            for el in paragraph.runs[0]._r.findall(f"{{{_A}}}t")
        ]
        self.assertEqual(texts, ["안녕", "세계"])
        self.assertEqual(
            len(paragraph.runs[0]._r.findall(f"{{{_A}}}br")), 1
        )

    def test_single_line_translation_drops_original_break(self):
        """Flattened (single-line) translations must not leak the old lines."""
        prs = Presentation()
        _, paragraph = _add_paragraph(prs)
        run = paragraph.add_run()
        run.text = "Hello"
        _append_break(run, "World")

        self._apply(paragraph, "Hello\nWorld", "안녕하세요 세계")

        xml = _run_xml(paragraph.runs[0])
        self.assertNotIn("Hello", xml)
        self.assertNotIn("World", xml)
        self.assertNotIn("<a:br", xml)
        self.assertEqual(paragraph.text, "안녕하세요 세계")


class WriterHyperlinkTestCase(unittest.TestCase):
    def _make_linked_paragraph(self):
        # NOTE: the linked run comes first — that is the exact shape that
        # triggered the old bug (single-group path wrote the whole
        # translation into the linked run).
        prs = Presentation()
        slide, paragraph = _add_paragraph(prs)
        linked = paragraph.add_run()
        linked.text = "여기"
        _append_hyperlink(linked, slide)
        after = paragraph.add_run()
        after.text = "를 클릭하면 안내로 이동합니다"
        return paragraph

    def test_linked_run_not_grouped_with_plain_text(self):
        paragraph = self._make_linked_paragraph()
        groups = _group_runs_by_format(list(paragraph.runs))
        self.assertEqual(len(groups), 2)

    def test_translation_does_not_spread_link_over_paragraph(self):
        paragraph = self._make_linked_paragraph()
        info = types.SimpleNamespace(
            paragraph=paragraph,
            original_text="여기를 클릭하면 안내로 이동합니다",
            is_note=False,
            slide_index=0,
            shape_index=0,
            paragraph_index=0,
        )
        translation = "Click here for details"
        PPTWriter().apply_translations(
            [info], [translation], Presentation(), text_fit_mode="none"
        )

        self.assertEqual(paragraph.text, translation)
        linked_runs = [
            run
            for run in paragraph.runs
            if run._r.find(f"{{{_A}}}hlinkClick") is not None
        ]
        # Exactly the original linked run keeps the hyperlink; the body
        # text lives in the plain runs (link dropped, text preserved).
        self.assertEqual(len(linked_runs), 1)
        self.assertNotEqual(linked_runs[0].text, translation)


if __name__ == "__main__":
    unittest.main()
